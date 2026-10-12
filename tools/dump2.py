import sys, io
from pptx import Presentation
paths = sys.argv[1:]
out = io.open("D:/project/tools/all_decks.txt", "w", encoding="utf-8")
def dump(path):
    prs = Presentation(path)
    out.write("="*90 + "\n")
    out.write("FILE: %s\nslides: %d\n" % (path, len(prs.slides)))
    out.write("="*90 + "\n")
    for i, slide in enumerate(prs.slides, 1):
        out.write("\n----- Slide %d -----\n" % i)
        for shape in slide.shapes:
            txt = ""
            if shape.has_text_frame:
                txt = " | ".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if shape.has_table:
                rows = []
                for r in shape.table.rows:
                    rows.append(" ; ".join(c.text.replace("\n"," ") for c in r.cells))
                txt = " || ".join(rows)
            if shape.has_chart:
                ch = shape.chart
                try: cats = list(ch.plots[0].categories)
                except Exception: cats = []
                series = ["%s=%s" % (s.name, list(s.values)) for s in ch.series]
                txt = "CHART %s cats=%s series=%s" % (ch.chart_type, cats, series)
            if txt.strip():
                out.write("  [%s] %s\n" % (shape.name, txt))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.write("  NOTES: %s\n" % slide.notes_slide.notes_text_frame.text.strip()[:800])
for p in paths:
    dump(p)
out.close()
print("ok")
