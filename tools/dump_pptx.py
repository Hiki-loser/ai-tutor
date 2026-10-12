import sys
from pptx import Presentation
from pptx.util import Emu

def dump(path):
    prs = Presentation(path)
    print("=" * 90)
    print("FILE:", path)
    print("slides:", len(prs.slides), " size:", prs.slide_width, "x", prs.slide_height)
    print("=" * 90)
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n----- Slide {i} -----")
        for shape in slide.shapes:
            t = shape.shape_type
            name = shape.name
            txt = ""
            if shape.has_text_frame:
                txt = " | ".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if shape.has_table:
                rows = []
                for r in shape.table.rows:
                    rows.append(" ; ".join(c.text.replace("\n", " ") for c in r.cells))
                txt = " || ".join(rows)
            if shape.has_chart:
                ch = shape.chart
                try:
                    cats = list(ch.plots[0].categories)
                except Exception:
                    cats = []
                series = []
                for s in ch.series:
                    series.append(f"{s.name}={list(s.values)}")
                txt = f"CHART {ch.chart_type} cats={cats} series={series}"
            if txt.strip():
                print(f"  [{name}] {txt}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print("  NOTES:", slide.notes_slide.notes_text_frame.text.strip()[:500])

for p in sys.argv[1:]:
    dump(p)
